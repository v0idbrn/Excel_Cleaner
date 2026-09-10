"""tests/test_analyzer_tolerant_headers.py — Regresión Paso 4 (plan Fiverr).

Cubre la tolerancia de encabezados verificada en la auditoría:
  ANTES: un CSV/XLSX con encabezado vacío era rechazado ENTERO
         (AnalyzerError 'encabezado vacío'), y una primera hoja vacía
         (portada corporativa) también rechazaba el archivo.
  DESPUÉS: la columna se auto-nombra 'Columna_N', el archivo se carga con
         todos sus datos y el Analyzer emite ENCABEZADO_AUTO_GENERADO para
         revisión. La primera hoja sin datos se salta hacia la primera hoja
         útil. Los duplicados REALES siguen rechazándose (cero pérdida).
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from analyzer import (
    AnalyzerError,
    analyze_file,
    build_file_info,
    load_dataframe,
)


class TestTolerantHeadersLoad(unittest.TestCase):
    """Motor de carga tolerante (el mismo que usan GUI y batch)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    # --- Encabezado vacío: ahora se carga, antes se rechazaba el archivo ---

    def test_csv_blank_header_loads_with_columna_n(self):
        p = self.tmp / "blank.csv"
        p.write_text("Nombre;;Ciudad\nAna;30;Salta\nLuis;25;CABA\n", encoding="utf-8")
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Columna_2", "Ciudad"])
        self.assertEqual(len(df), 2)

    def test_csv_whitespace_header_loads_with_columna_n(self):
        p = self.tmp / "ws.csv"
        p.write_text("Nombre, ,Ciudad\nAna,30,Salta\n", encoding="utf-8")
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Columna_2", "Ciudad"])

    def test_xlsx_blank_header_loads_with_columna_n(self):
        p = self.tmp / "blank.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Nombre", None, "Ciudad"])
        ws.append(["Ana", 30, "Salta"])
        wb.save(p)
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Columna_2", "Ciudad"])
        self.assertEqual(df.iloc[0, 1], 30)

    # --- Duplicados reales: siguen rechazados (no se adivina) ---

    def test_csv_duplicate_header_still_rejected(self):
        p = self.tmp / "dup.csv"
        p.write_text("ID,Nombre,Nombre\n1,A,X\n2,B,Y\n", encoding="utf-8")
        with self.assertRaises(AnalyzerError):
            load_dataframe(build_file_info(p))

    def test_xlsx_duplicate_header_still_rejected(self):
        p = self.tmp / "dup.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Nombre", "Nombre"])
        ws.append([1, "A", "X"])
        wb.save(p)
        with self.assertRaises(AnalyzerError):
            load_dataframe(build_file_info(p))

    def test_duplicate_insensitive_to_case_or_spaces_rejected(self):
        # 'Nombre' y ' NOMBRE ' colisionan tras normalizar: pandas las manglearía
        # silenciosamente ('Nombre.1'); se rechaza para forzar la corrección.
        p = self.tmp / "dup_loose.csv"
        p.write_text("ID,Nombre, NOMBRE \n1,A,X\n", encoding="utf-8")
        with self.assertRaises(AnalyzerError):
            load_dataframe(build_file_info(p))

    # --- Hojas vacías: se salta la portada, antes se rechazaba el archivo ---

    def test_empty_first_sheet_skipped_to_data_sheet(self):
        p = self.tmp / "empty_first.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Portada"
        ws2 = wb.create_sheet("Datos")
        ws2.append(["Cliente", "Monto"])
        ws2.append(["ACME", 100])
        wb.save(p)
        fi = build_file_info(p)
        df = load_dataframe(fi)
        self.assertEqual(list(df.columns), ["Cliente", "Monto"])
        self.assertEqual(df.iloc[0, 0], "ACME")

    def test_all_sheets_empty_rejected_clearly(self):
        p = self.tmp / "all_empty.xlsx"
        wb = openpyxl.Workbook()
        wb.create_sheet("A")
        wb.create_sheet("B")
        wb.save(p)
        with self.assertRaises(AnalyzerError) as ctx:
            load_dataframe(build_file_info(p))
        self.assertIn("hojas con datos", str(ctx.exception))

    def test_xlsx_headers_only_sheet_loads_zero_rows(self):
        p = self.tmp / "headers_only.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["ColA", "ColB"])
        wb.save(p)
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["ColA", "ColB"])
        self.assertEqual(len(df), 0)

    def test_xlsx_ragged_row_unnamed_column_renamed(self):
        # Fila de datos más larga que el encabezado: pandas crea 'Unnamed: N'.
        p = self.tmp / "ragged.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A", "B"])
        ws.append([1, 2, 3])
        wb.save(p)
        df = load_dataframe(build_file_info(p))
        self.assertIn("Columna_3", list(df.columns))
        self.assertNotIn("Unnamed: 2", list(df.columns))

    # --- El caso feliz no cambia ---

    def test_clean_csv_unaffected_by_tolerance(self):
        p = self.tmp / "clean.csv"
        p.write_text("Nombre,Edad,Ciudad\nAna,30,Salta\n", encoding="utf-8")
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Edad", "Ciudad"])
        self.assertEqual(len(df), 1)


class TestUnnamedColumnIssue(unittest.TestCase):
    """El Analyzer emite ENCABEZADO_AUTO_GENERADO para las columnas auto-nombradas."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_issue_emitted_for_auto_named_column(self):
        p = self.tmp / "blank.csv"
        p.write_text("Nombre;;Edad\nAna;;30\nLuis;;25\n", encoding="utf-8")
        report = analyze_file(p)
        issues = [i for i in report.issues if i.category == "ENCABEZADO_AUTO_GENERADO"]
        self.assertEqual(len(issues), 1)
        issue = issues[0]
        self.assertEqual(issue.column, "Columna_2")
        self.assertEqual(issue.suggested_action, "revisar_manualmente")
        self.assertEqual(issue.severity.value, "MEDIA")
        self.assertIn("Columna_2", issue.description)

    def test_no_issue_on_clean_file(self):
        p = self.tmp / "clean.csv"
        p.write_text("Nombre,Edad\nAna,30\n", encoding="utf-8")
        report = analyze_file(p)
        self.assertFalse(
            [i for i in report.issues if i.category == "ENCABEZADO_AUTO_GENERADO"]
        )

    def test_unnamed_issue_with_data_preserved(self):
        # La columna auto-nombrada conserva sus datos (affected_count = no nulos).
        p = self.tmp / "blank.csv"
        p.write_text("Nombre;;Edad\nAna;X;30\nLuis;Y;25\n", encoding="utf-8")
        report = analyze_file(p)
        issue = next(i for i in report.issues if i.category == "ENCABEZADO_AUTO_GENERADO")
        self.assertEqual(issue.affected_count, 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
