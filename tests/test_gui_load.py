"""tests/test_gui_load.py — Regresión Paso 1 (plan Fiverr).

Cubre dos defectos verificados en la auditoría:
  P0-3: gui/app.py usaba process_batch_files sin importarlo (NameError en batch multi-archivo).
  P0-4: _worker_load cargaba con pandas crudo -> CSV ';' roto (1 columna) y
        XLSX multi-hoja sin selector (solo hoja 1).

La carga corregida pasa por el Analyzer (build_file_info + load_dataframe), el mismo
motor testado que usa el batch, con selector de hoja para XLSX multi-hoja.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from analyzer import AnalyzerError, build_file_info, load_dataframe


class TestAnalyzerLoadEngine(unittest.TestCase):
    """El motor de carga que ahora usa la GUI (mismo que batch/CLI)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_semicolon_csv_loads_three_columns(self):
        # ANTES (pandas crudo): 1 columna 'Nombre;Edad;Ciudad'.
        p = self.tmp / "semi.csv"
        p.write_text("Nombre;Edad;Ciudad\nAna;30;Salta\nLuis;25;CABA\n", encoding="utf-8")
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Edad", "Ciudad"])
        self.assertEqual(len(df), 2)

    def test_utf8_bom_headers_clean(self):
        p = self.tmp / "bom.csv"
        p.write_bytes("﻿Nombre,Edad\nAna,30\n".encode("utf-8"))
        df = load_dataframe(build_file_info(p))
        self.assertEqual(list(df.columns), ["Nombre", "Edad"])

    def test_empty_file_rejected(self):
        p = self.tmp / "vacio.csv"
        p.write_bytes(b"")
        with self.assertRaises(AnalyzerError):
            build_file_info(p)

    def test_multi_sheet_lists_all_sheets_and_loads_chosen(self):
        p = self.tmp / "multi.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Ventas"
        wb.active.append(["Prod", "Monto"])
        wb.active.append(["A", 10])
        ws2 = wb.create_sheet("Clientes")
        ws2.append(["Cliente"])
        ws2.append(["Bob"])
        wb.save(p)

        fi = build_file_info(p)
        self.assertEqual(fi.sheet_names, ("Ventas", "Clientes"))

        df = load_dataframe(fi, "Clientes")
        self.assertEqual(list(df.columns), ["Cliente"])
        self.assertEqual(df.iloc[0, 0], "Bob")


class TestGuiLoadWorker(unittest.TestCase):
    """El worker de la GUI carga vía Analyzer; los errores van a la cola, no crashean."""

    @classmethod
    def setUpClass(cls):
        try:
            import tkinter as tk
            cls.root = tk.Tk()
            cls.root.withdraw()
        except Exception:
            raise unittest.SkipTest("Tkinter no disponible en este entorno")

    @classmethod
    def tearDownClass(cls):
        try:
            cls.root.destroy()
        except Exception:
            pass

    def setUp(self):
        from gui.app import ExcelCleanerApp
        self.app = ExcelCleanerApp(self.root)
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        self.app.reset_state()

    def _drain_queue(self):
        import queue
        msgs = []
        try:
            while True:
                msgs.append(self.app.task_queue.get_nowait())
        except queue.Empty:
            pass
        return msgs

    def test_worker_load_semicolon_csv_three_columns(self):
        p = self.tmp / "semi.csv"
        p.write_text("Nombre;Edad\nAna;30\nLuis;25\n", encoding="utf-8")
        self.app._worker_load(str(p))
        msgs = self._drain_queue()
        self.assertTrue(any(t == "LOAD_DONE" for t, _ in msgs), f"mensajes: {msgs}")
        self.assertFalse(any(t == "ERROR" for t, _ in msgs))
        self.assertEqual(list(self.app.original_df.columns), ["Nombre", "Edad"])

    def test_worker_load_multi_sheet_xlsx_chooses_sheet(self):
        p = self.tmp / "multi.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Ventas"
        wb.active.append(["Prod", "Monto"])
        ws2 = wb.create_sheet("Clientes")
        ws2.append(["Cliente"])
        ws2.append(["Bob"])
        wb.save(p)
        self.app._worker_load(str(p), "Clientes")
        msgs = self._drain_queue()
        self.assertTrue(any(t == "LOAD_DONE" for t, _ in msgs), f"mensajes: {msgs}")
        self.assertEqual(list(self.app.original_df.columns), ["Cliente"])

    def test_worker_load_invalid_file_posts_error_not_crash(self):
        p = self.tmp / "vacio.csv"
        p.write_bytes(b"")
        self.app._worker_load(str(p))
        msgs = self._drain_queue()
        self.assertTrue(any(t == "ERROR" for t, _ in msgs), f"mensajes: {msgs}")
        self.assertIsNone(self.app.original_df)

    def test_worker_load_oversize_rejected_by_size_limit(self):
        # build_file_info aplica MAX_FILE_SIZE_BYTES antes de leer (regla de la auditoría).
        import analyzer, unittest.mock as mock
        p = self.tmp / "grande.csv"
        p.write_text("A\n1\n", encoding="utf-8")
        with mock.patch.object(analyzer.config, "MAX_FILE_SIZE_BYTES", 1):
            self.app._worker_load(str(p))
        msgs = self._drain_queue()
        self.assertTrue(any(t == "ERROR" for t, _ in msgs), f"mensajes: {msgs}")


class TestBatchMultiSelectImport(unittest.TestCase):
    """P0-3: process_batch_files debe estar importado y resoluble en gui.app."""

    def test_process_batch_files_resolvable_from_gui_app(self):
        import gui.app as gui_app
        self.assertTrue(callable(getattr(gui_app, "process_batch_files", None)))
        from batch_processor import process_batch_files as real
        self.assertIs(gui_app.process_batch_files, real)


if __name__ == "__main__":
    unittest.main(verbosity=2)
